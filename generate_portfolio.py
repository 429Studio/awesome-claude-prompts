"""
Event Photo Portfolio - PPTX Generator
Minimal & Elegant style, 10 slides
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── Palette ──────────────────────────────────────────────
BG_DARK   = RGBColor(0x1A, 0x1A, 0x1A)
BG_MED    = RGBColor(0x22, 0x22, 0x22)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
GREY      = RGBColor(0xAA, 0xAA, 0xAA)
ACCENT    = RGBColor(0xC8, 0xA2, 0x6E)  # warm gold

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H

# Use blank layout
blank_layout = prs.slide_layouts[6]


def set_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, text, left, top, width, height,
             font_size=18, color=WHITE, bold=False, alignment=PP_ALIGN.LEFT,
             font_name="Calibri Light"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return tf


def add_divider(slide, left, top, width, color=ACCENT):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left, top, width, Pt(2)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_photo_placeholder(slide, left, top, width, height, label="[Your Photo Here]"):
    """Adds a styled rectangle placeholder for a photo."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0x2A, 0x2A, 0x2A)
    shape.line.color.rgb = RGBColor(0x44, 0x44, 0x44)
    shape.line.width = Pt(1)
    # Label inside
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = label
    p.font.size = Pt(11)
    p.font.color.rgb = GREY
    p.font.name = "Calibri Light"
    p.alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].space_before = Pt(0)
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER


# ═══════════════════════════════════════════════════════════
# SLIDE 1 — Cover
# ═══════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, BG_DARK)

add_text(slide, "EVENT PHOTOGRAPHY",
         Inches(1.5), Inches(1.8), Inches(10), Inches(1),
         font_size=48, color=ACCENT, bold=True, alignment=PP_ALIGN.CENTER,
         font_name="Calibri")

add_divider(slide, Inches(5.5), Inches(2.95), Inches(2.3))

add_text(slide, "PORTFOLIO",
         Inches(1.5), Inches(3.2), Inches(10), Inches(0.8),
         font_size=32, color=WHITE, bold=False, alignment=PP_ALIGN.CENTER)

add_text(slide, "[Your Name / Studio Name]",
         Inches(1.5), Inches(4.3), Inches(10), Inches(0.6),
         font_size=16, color=GREY, alignment=PP_ALIGN.CENTER)

add_text(slide, "Capturing Moments That Matter",
         Inches(1.5), Inches(5.0), Inches(10), Inches(0.5),
         font_size=14, color=GREY, alignment=PP_ALIGN.CENTER,
         font_name="Calibri Light")


# ═══════════════════════════════════════════════════════════
# SLIDE 2 — About Me
# ═══════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, BG_DARK)

add_text(slide, "ABOUT ME",
         Inches(0.8), Inches(0.5), Inches(5), Inches(0.7),
         font_size=28, color=ACCENT, bold=True)

add_divider(slide, Inches(0.8), Inches(1.25), Inches(1.5))

about_text = (
    "I am a professional event photographer with a passion for capturing "
    "authentic moments across all types of events — from intimate weddings "
    "to large-scale corporate conferences, live concerts, and everything "
    "in between.\n\n"
    "My approach blends documentary storytelling with an eye for detail, "
    "ensuring every event is remembered through genuine, emotive imagery.\n\n"
    "Based in [Your City]. Available worldwide."
)
add_text(slide, about_text,
         Inches(0.8), Inches(1.6), Inches(5.5), Inches(4.5),
         font_size=14, color=WHITE)

# Profile photo placeholder
add_photo_placeholder(slide, Inches(7.8), Inches(1.2), Inches(4.5), Inches(5),
                      "[Profile / Hero Photo]")


# ═══════════════════════════════════════════════════════════
# SLIDE 3 — Services Overview
# ═══════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, BG_DARK)

add_text(slide, "WHAT I COVER",
         Inches(0.8), Inches(0.5), Inches(10), Inches(0.7),
         font_size=28, color=ACCENT, bold=True, alignment=PP_ALIGN.CENTER)

add_divider(slide, Inches(5.8), Inches(1.25), Inches(1.5))

services = [
    ("Weddings & Social Events", "Engagements, receptions,\nbirthdays, anniversaries"),
    ("Corporate & Conferences", "Summits, product launches,\ntrade shows, galas"),
    ("Concerts & Entertainment", "Live music, festivals,\ntheater, nightlife events"),
    ("Personal & Lifestyle", "Portraits, brand shoots,\nspecial occasions"),
]

for i, (title, desc) in enumerate(services):
    col = i % 4
    x = Inches(0.6 + col * 3.15)
    y = Inches(2.0)

    # Icon placeholder
    add_photo_placeholder(slide, x, y, Inches(2.8), Inches(2.8),
                          f"[{title} Photo]")

    add_text(slide, title,
             x, y + Inches(3.0), Inches(2.8), Inches(0.5),
             font_size=13, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

    add_text(slide, desc,
             x, y + Inches(3.4), Inches(2.8), Inches(0.8),
             font_size=11, color=GREY, alignment=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════
# SLIDE 4-7 — Gallery Spreads (4 slides, 2-3 photos each)
# ═══════════════════════════════════════════════════════════
gallery_slides = [
    ("WEDDINGS & SOCIAL", [
        "[Wedding Photo 1]", "[Wedding Photo 2]", "[Wedding Photo 3]"
    ]),
    ("CORPORATE & CONFERENCES", [
        "[Corporate Photo 1]", "[Corporate Photo 2]", "[Corporate Photo 3]"
    ]),
    ("CONCERTS & ENTERTAINMENT", [
        "[Concert Photo 1]", "[Concert Photo 2]", "[Concert Photo 3]"
    ]),
    ("LIFESTYLE & PORTRAITS", [
        "[Lifestyle Photo 1]", "[Lifestyle Photo 2]", "[Lifestyle Photo 3]"
    ]),
]

for title, photos in gallery_slides:
    slide = prs.slides.add_slide(blank_layout)
    set_bg(slide, BG_DARK)

    add_text(slide, title,
             Inches(0.8), Inches(0.4), Inches(11), Inches(0.6),
             font_size=22, color=ACCENT, bold=True)

    add_divider(slide, Inches(0.8), Inches(1.05), Inches(1.2))

    # 3-photo layout: one large left, two stacked right
    add_photo_placeholder(slide, Inches(0.8), Inches(1.4), Inches(7), Inches(5.5),
                          photos[0])
    add_photo_placeholder(slide, Inches(8.1), Inches(1.4), Inches(4.4), Inches(2.6),
                          photos[1])
    add_photo_placeholder(slide, Inches(8.1), Inches(4.3), Inches(4.4), Inches(2.6),
                          photos[2])


# ═══════════════════════════════════════════════════════════
# SLIDE 8 — Key Numbers / Stats
# ═══════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, BG_MED)

add_text(slide, "BY THE NUMBERS",
         Inches(0.5), Inches(0.5), Inches(12), Inches(0.7),
         font_size=28, color=ACCENT, bold=True, alignment=PP_ALIGN.CENTER)

add_divider(slide, Inches(5.8), Inches(1.25), Inches(1.5))

stats = [
    ("500+", "Events Covered"),
    ("50K+", "Photos Delivered"),
    ("200+", "Happy Clients"),
    ("8+", "Years of Experience"),
]

for i, (number, label) in enumerate(stats):
    x = Inches(0.8 + i * 3.15)
    y = Inches(2.5)

    add_text(slide, number,
             x, y, Inches(2.8), Inches(1.2),
             font_size=48, color=ACCENT, bold=True, alignment=PP_ALIGN.CENTER,
             font_name="Calibri")

    add_text(slide, label,
             x, y + Inches(1.2), Inches(2.8), Inches(0.5),
             font_size=16, color=WHITE, alignment=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════
# SLIDE 9 — Testimonial
# ═══════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, BG_DARK)

add_text(slide, "\u201c",
         Inches(1), Inches(1), Inches(1), Inches(1.5),
         font_size=96, color=ACCENT, bold=True)

testimonial = (
    "Absolutely phenomenal work. Every single photo captured the energy "
    "and emotion of our event perfectly. We couldn't have asked for a "
    "better photographer. Highly recommend!"
)
add_text(slide, testimonial,
         Inches(2), Inches(2.2), Inches(9), Inches(2.5),
         font_size=20, color=WHITE, alignment=PP_ALIGN.CENTER,
         font_name="Calibri Light")

add_text(slide, "\u2014 [Client Name], [Event Name]",
         Inches(2), Inches(4.8), Inches(9), Inches(0.5),
         font_size=14, color=GREY, alignment=PP_ALIGN.CENTER)

add_text(slide, "[Add more testimonials as needed]",
         Inches(2), Inches(5.8), Inches(9), Inches(0.4),
         font_size=11, color=RGBColor(0x55, 0x55, 0x55), alignment=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════
# SLIDE 10 — Contact / CTA
# ═══════════════════════════════════════════════════════════
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, BG_DARK)

add_text(slide, "LET'S WORK TOGETHER",
         Inches(1), Inches(1.5), Inches(11), Inches(0.8),
         font_size=36, color=ACCENT, bold=True, alignment=PP_ALIGN.CENTER,
         font_name="Calibri")

add_divider(slide, Inches(5.5), Inches(2.45), Inches(2.3))

contact_info = (
    "[Your Name]\n\n"
    "Email: [your@email.com]\n"
    "Phone: [+1 234 567 8900]\n"
    "Website: [www.yourwebsite.com]\n"
    "Instagram: [@yourhandle]"
)
add_text(slide, contact_info,
         Inches(1), Inches(3.0), Inches(11), Inches(3.5),
         font_size=16, color=WHITE, alignment=PP_ALIGN.CENTER)

add_text(slide, "Available for bookings worldwide  |  Inquire for custom packages",
         Inches(1), Inches(6.2), Inches(11), Inches(0.5),
         font_size=12, color=GREY, alignment=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════
# Save
# ═══════════════════════════════════════════════════════════
output_path = "/home/user/awesome-claude-prompts/Event_Photo_Portfolio.pptx"
prs.save(output_path)
print(f"Portfolio saved to: {output_path}")
print(f"Total slides: {len(prs.slides)}")
